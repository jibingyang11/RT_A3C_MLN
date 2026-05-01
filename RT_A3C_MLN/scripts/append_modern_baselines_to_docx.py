from __future__ import annotations

from pathlib import Path
import shutil

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
ARTIFACTS = ROOT / "paper_artifacts"
MODERN = ARTIFACTS / "modern_baselines"
TABLES = MODERN / "tables"
FIGURES = MODERN / "figures"
DOCX = ARTIFACTS / "RT_A3C_MLN_实验章节报告.docx"

ACCENT = "1F4E79"
HEADER_FILL = "D9EAF7"
NOTE_FILL = "EAF3F8"


def set_east_asian_font(run, font="宋体", size=None, bold=None, color=None):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=55, bottom=70, end=55):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="B7C7D6", size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_section_portrait(section):
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)


def set_section_landscape(section):
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21)
    section.top_margin = Cm(1.45)
    section.bottom_margin = Cm(1.35)
    section.left_margin = Cm(1.15)
    section.right_margin = Cm(1.15)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_east_asian_font(r, font="微软雅黑", size=15 if level == 1 else 12.5, bold=True, color=ACCENT)
    return p


def add_body(doc, text, first_line=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(5)
    if first_line:
        p.paragraph_format.first_line_indent = Pt(21)
    r = p.add_run(text)
    set_east_asian_font(r, font="宋体", size=10.5)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(text)
    set_east_asian_font(r, font="宋体", size=9.2, bold=True, color="333333")
    return p


def add_note_box(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_borders(table, color="BFD3E3", size="6")
    cell = table.cell(0, 0)
    shade_cell(cell, NOTE_FILL)
    set_cell_margins(cell, top=150, bottom=150, start=180, end=180)
    p = cell.paragraphs[0]
    r = p.add_run(title)
    set_east_asian_font(r, font="微软雅黑", size=10.5, bold=True, color=ACCENT)
    p2 = cell.add_paragraph()
    p2.paragraph_format.line_spacing = 1.18
    r2 = p2.add_run(body)
    set_east_asian_font(r2, font="宋体", size=10)
    doc.add_paragraph()


def format_word_table(table, font_size=6.2, header_size=6.0):
    table.style = "Table Grid"
    table.autofit = True
    set_table_borders(table)
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=60, bottom=60, start=45, end=45)
            if i == 0:
                shade_cell(cell, HEADER_FILL)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    set_east_asian_font(
                        r,
                        font="宋体",
                        size=header_size if i == 0 else font_size,
                        bold=True if i == 0 else None,
                    )
    set_repeat_table_header(table.rows[0])


def prepare_table5():
    df = pd.read_csv(TABLES / "Table5_Modern_Stream_Baselines.csv")
    rename = {
        "Method": "方法",
        "Lineage/year": "年份/谱系",
        "Rule-structure adaptation": "规则结构自适应",
        "Core F1": "常规F1",
        "Core Log-loss": "常规Log-loss",
        "Core cycle ms": "常规周期(ms)",
        "Core 10ms %": "10ms成功率(%)",
        "Strong post-F1": "强漂移F1",
        "Strong post Log-loss": "强漂移Log-loss",
        "Recovery lag": "恢复延迟",
        "Struct update %": "结构更新(%)",
        "Active emerging rules": "激活新规则",
    }
    df = df.rename(columns=rename)
    df["规则结构自适应"] = df["规则结构自适应"].replace({"Yes": "是", "No": "否"})
    df["年份/谱系"] = df["年份/谱系"].replace(
        {
            "2026 / ours": "2026/本文",
            "fixed-rule online reference": "固定规则在线参考",
            "static reference": "静态参考",
            "2021 framework": "River 2021",
            "2009 stream tree": "HAT 2009",
            "2017 stream ensemble": "ARF 2017",
            "2019 stream ensemble": "SRP 2019",
        }
    )
    return df


def add_dataframe_table(doc, df, caption):
    add_caption(doc, caption)
    table = doc.add_table(rows=1, cols=len(df.columns))
    for j, col in enumerate(df.columns):
        table.cell(0, j).text = str(col)
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for j, col in enumerate(df.columns):
            cells[j].text = str(row[col])
    format_word_table(table)
    doc.add_paragraph()


def add_picture(doc, image, caption, width_inches=6.35):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image), width=Inches(width_inches))
    add_caption(doc, caption)


def append_section(doc):
    table5 = prepare_table5()

    doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_portrait(doc.sections[-1])
    add_heading(doc, "4.11 现代数据流基线补充实验", 1)
    add_body(
        doc,
        "为回应对比方法年份偏早的问题，本文在原有经典 MLN 基线之外，进一步加入数据流学习领域常用的现代或近现代在线基线。新增基线包括 River 框架中的在线逻辑回归、Hoeffding Adaptive Tree（HAT）、Adaptive Random Forest（ARF）和 Streaming Random Patches（SRP）。这些方法并非 MLN 结构学习方法，但代表了数据流分类和概念漂移处理中常见的在线更新范式，可用于检验本文方法的实时响应能力是否仅来自简单在线分类器，而不是规则结构自适应机制。",
    )
    add_body(
        doc,
        "补充实验仍采用两类场景：第一类为四个真实数据集上的常规数据流实验，用于比较平均 F1、Log-loss、阻塞周期和 10 ms 截止成功率；第二类为强规则漂移压力测试，用于观察当后半段数据流出现初始阶段几乎不可见的新判别规则时，各方法能否恢复性能。需要强调的是，HAT、ARF、SRP 等方法可以更新预测器参数或树结构，但不产生 MLN 规则的新增、删除和权重解释，因此它们只能作为数据流学习基线，不能替代概率逻辑规则结构自适应方法。",
    )

    doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_landscape(doc.sections[-1])
    add_dataframe_table(doc, table5, "表4-5 现代数据流基线补充实验结果")
    add_body(
        doc,
        "表4-5 显示，在常规数据流实验中，RT-A3C-MLN 的平均阻塞周期为 2.752 ms，10 ms 成功率为 100%，满足实时推理约束。OnlineWeight MLN 和 Static MaxEnt 的周期更低，但二者均不支持规则结构自适应；当进入强规则漂移场景后，其漂移后 F1 分别下降到 0.416 和 0.217，且均未恢复到设定阈值。该结果说明，固定规则结构方法在温和数据流中可能速度很快，但不能证明其具备规则层面的自适应能力。",
    )
    add_body(
        doc,
        "River Online LR、HAT、ARF 和 SRP 代表流式学习基线。它们在强漂移场景中可以通过在线学习或集成更新获得一定分类性能，但其阻塞周期明显高于 10 ms 阈值：River Online LR、HAT、ARF 和 SRP 的常规阻塞周期分别为 143.948 ms、578.718 ms、130.601 ms 和 3598.283 ms，10 ms 成功率均为 0%。同时，这些方法的结构更新比例和激活新规则数均为 0，因为它们不会执行 MLN 规则新增、删除或规则权重结构解释。相比之下，RT-A3C-MLN 在强漂移后激活的新规则数为 5.70，结构更新比例为 90.0%，并在 3 个 chunk 后恢复，直接支撑本文关于“规则可在数据流输入下自适应调整”的核心论点。",
    )

    doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_portrait(doc.sections[-1])
    add_picture(
        doc,
        FIGURES / "Fig9_Modern_Stream_Baselines.png",
        "图4-9 现代数据流基线的实时性、漂移恢复和规则自适应对比",
        width_inches=6.35,
    )
    add_body(
        doc,
        "图4-9 左侧子图给出常规数据流中的阻塞周期，纵轴采用对数坐标，红色虚线表示 10 ms 实时阈值。RT-A3C-MLN 位于阈值以下，而 River Online LR、HAT、ARF 和 SRP 均明显超过阈值，说明普通数据流算法虽然支持在线更新，但逐样本更新和集成维护带来的阻塞代价并不一定适合本文设定的实时概率推理场景。",
    )
    add_body(
        doc,
        "中间子图同时比较常规数据流 F1 和强规则漂移后的 post-F1。可以看到，一些流式分类器在强漂移后仍具有一定分类能力，但这类性能来自分类器参数或树节点更新，并不对应 MLN 规则结构的可解释变化。右侧子图进一步显示，只有 RT-A3C-MLN 在强漂移后产生非零的 active emerging rules，同时恢复延迟较短。这说明本文方法的优势不是把 MLN 换成普通在线分类器，而是在保持低阻塞周期的同时，保留了概率逻辑规则网络的结构更新证据。",
    )
    add_note_box(
        doc,
        "现代基线补充实验结论",
        "新增数据流基线说明：即使与 HAT、ARF、SRP 等在线学习方法比较，RT-A3C-MLN 仍能在实时截止约束内完成推理和更新，并在强规则漂移下产生明确的规则结构变化。该补充实验可用于论文中回应“经典 MLN 基线较老”的疑问，同时强化本文主线：本文方法不是追求静态准确率最高，而是证明 MLN 能够在数据流下快速、自适应、可解释地调整规则并实时输出概率推理结果。",
    )


def main():
    if not DOCX.exists():
        raise FileNotFoundError(DOCX)
    doc = Document(str(DOCX))
    if any("现代数据流基线补充实验" in paragraph.text for paragraph in doc.paragraphs):
        raise RuntimeError("目标文档中已存在现代数据流基线补充实验小节，已停止以避免重复追加。")
    backup = DOCX.with_name(DOCX.stem + "_backup_before_Table5_Fig9.docx")
    if not backup.exists():
        shutil.copy2(DOCX, backup)
    append_section(doc)
    doc.save(str(DOCX))
    print(DOCX)


if __name__ == "__main__":
    main()
