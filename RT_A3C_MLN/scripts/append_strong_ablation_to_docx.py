from __future__ import annotations

from pathlib import Path
import shutil
import sys

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.shared import Inches

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts.append_strong_rule_drift_to_docx import (
    ROOT,
    add_body,
    add_caption,
    add_heading,
    add_note_box,
    format_word_table,
    set_section_landscape,
    set_section_portrait,
    target_docx,
)


ARTIFACTS = ROOT / "paper_artifacts"
ABLATION = ARTIFACTS / "strong_rule_drift_ablation"
FIGURES = ABLATION / "figures"
TABLES = ABLATION / "tables"


def prepare_table4():
    df = pd.read_csv(TABLES / "Table4_Strong_Drift_Ablation.csv")
    delta_col = next((col for col in df.columns if "F1 vs Full" in col), "ΔF1 vs Full")
    rename = {
        "Variant": "版本",
        "Controlled change": "控制变量",
        "Post-drift F1": "漂移后F1",
        delta_col: "相对Full F1差",
        "Post-drift Log-loss": "漂移后Log-loss",
        "Recovery lag": "恢复延迟",
        "Blocking cycle ms": "阻塞周期(ms)",
        "10ms success %": "10ms成功率(%)",
        "Struct update %": "结构更新(%)",
        "Rule +/- per chunk": "规则增/删",
        "Active emerging rules": "激活新规则",
    }
    df = df.rename(columns=rename)
    df["版本"] = df["版本"].replace(
        {
            "Full RT-A3C": "Full RT-A3C",
            "No structure": "No structure",
            "No prior-shift": "No prior-shift",
            "Online only": "Online only",
        }
    )
    df["控制变量"] = df["控制变量"].replace(
        {
            "complete model": "完整模型",
            "disables background rule-structure refresh": "关闭后台结构刷新",
            "removes local prior-shift adaptation": "关闭局部先验漂移",
            "removes A3C structure learning": "移除A3C结构学习",
        }
    )
    return df


def add_dataframe_table(doc, df, caption):
    add_caption(doc, caption)
    table = doc.add_table(rows=1, cols=len(df.columns))
    for j, col in enumerate(df.columns):
        table.cell(0, j).text = str(col)
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for j, col in enumerate(df.columns):
            cells[j].text = str(row[col])
    format_word_table(table, font_size=7.2, header_size=7.0)
    doc.add_paragraph()


def add_picture(doc, image, caption, width_inches=6.25):
    p = doc.add_paragraph()
    p.alignment = 1
    run = p.add_run()
    run.add_picture(str(image), width=Inches(width_inches))
    add_caption(doc, caption)


def append_section(doc):
    table4 = prepare_table4()

    doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_portrait(doc.sections[-1])
    add_heading(doc, "4.10 强规则漂移下的结构刷新消融实验", 1)
    add_body(
        doc,
        "为了进一步证明后台规则结构刷新在强规则漂移场景中的必要性，本节在同一强规则漂移数据流上增加结构刷新消融实验。该实验保持数据流、漂移点、A3C 训练参数、前台在线权重更新和局部校准机制一致，仅关闭后台规则结构刷新，得到 No structure refresh 版本。这样可以直接检验：当后半段出现初始阶段几乎不可见的新判别规则时，仅依赖已有规则权重更新是否足以恢复模型性能。",
    )
    add_body(
        doc,
        "实验对比四个版本：Full RT-A3C 为完整模型；No structure 关闭后台规则结构刷新；No prior-shift 关闭局部先验漂移更新，用于区分概率校准与结构刷新作用；Online only 移除 A3C 结构学习，仅保留在线权重更新。评价指标重点关注漂移后的 F1、Log-loss、恢复延迟、结构更新比例、规则增删数量以及激活的新规则数量。",
    )

    doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_landscape(doc.sections[-1])
    add_dataframe_table(doc, table4, "表4-4 强规则漂移下的结构刷新消融实验")
    add_body(
        doc,
        "表4-4 表明，完整 Full RT-A3C 在规则漂移后的 F1 为 0.704，Log-loss 为 0.920，恢复延迟为 2 个 chunk，且结构更新比例达到 60.0%。相比之下，No structure 的漂移后 F1 下降到 0.340，相对完整模型下降 0.363，Log-loss 增加到 1.679，并且未能恢复到设定阈值。更关键的是，No structure 的结构更新比例、规则新增和规则删除均为 0，激活的新规则数量仅为 2.00，说明关闭后台结构刷新后，模型无法充分吸收后半段才出现的 emerging rules。",
    )
    add_body(
        doc,
        "No prior-shift 的漂移后 F1 与 Full 基本一致，结构更新比例和规则增删数量也保持一致，说明在该强规则漂移场景下，性能恢复主要来自规则结构刷新，而不是单纯的先验漂移校准。Online only 的 F1 为 0.347，Log-loss 达到 7.542，进一步说明仅依靠在线权重更新无法替代 A3C 结构学习。该消融结果直接证明：后台规则结构刷新不是辅助性模块，而是 RT-A3C-MLN 在强规则漂移下恢复性能的关键机制。",
    )

    doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_portrait(doc.sections[-1])
    add_picture(
        doc,
        FIGURES / "Fig8_Strong_Drift_Structure_Ablation.png",
        "图4-8 强规则漂移下的结构刷新消融结果",
        width_inches=6.25,
    )
    add_body(
        doc,
        "图4-8 从三个方面可视化结构刷新消融结果。左侧子图显示漂移后 F1，Full RT-A3C 和 No prior-shift 均保持在 0.704 左右，而 No structure 和 Online only 分别只有 0.340 和 0.347，说明一旦关闭结构刷新或移除 A3C 结构学习，模型难以适应后半段新出现的规则。中间子图显示漂移后 Log-loss，Online only 的概率损失最高，No structure 也明显高于 Full，说明固定已有规则上的权重调整会导致概率推理质量下降。",
    )
    add_body(
        doc,
        "右侧子图比较 Full RT-A3C 与 No structure 在数据流过程中的 active emerging rules 数量。规则漂移点后，Full RT-A3C 激活的新规则数量从约 4 个提升到 6 至 7 个，而 No structure 始终停留在约 2 个，几乎没有吸收新的判别规则。这一结果说明，性能差距不是由训练随机波动造成，而是由是否允许规则结构新增和替换导致。",
    )
    add_note_box(
        doc,
        "结构刷新消融结论",
        "强规则漂移压力测试下，No structure refresh 的漂移后 F1 明显低于 Full RT-A3C，且无法恢复；同时其结构更新比例、规则增删和激活新规则数量均显著不足。因此可以在论文中明确写出：后台规则结构刷新是本文方法实现数据流规则自适应的必要组件，仅靠固定规则上的在线权重更新不能支撑强规则漂移场景下的实时概率知识网络。",
    )


def main():
    doc_path = target_docx()
    doc = Document(str(doc_path))
    if any("强规则漂移下的结构刷新消融实验" in p.text for p in doc.paragraphs):
        raise RuntimeError("目标文档中已存在结构刷新消融小节，为避免重复追加已停止。")

    backup = doc_path.with_name(doc_path.stem + "_backup_before_Table4_Fig8.docx")
    if not backup.exists():
        shutil.copy2(doc_path, backup)

    append_section(doc)
    doc.save(str(doc_path))
    print(doc_path)


if __name__ == "__main__":
    main()
