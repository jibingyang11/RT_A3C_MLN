from __future__ import annotations

from pathlib import Path
import shutil

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
ARTIFACTS = ROOT / "paper_artifacts"
STRONG = ARTIFACTS / "strong_rule_drift"
FIGURES = STRONG / "figures"
TABLES = STRONG / "tables"

ACCENT = "1F4E79"
HEADER_FILL = "D9EAF7"
NOTE_FILL = "EAF3F8"


def target_docx() -> Path:
    matches = sorted(ARTIFACTS.glob("RT_A3C_MLN_*.docx"))
    if not matches:
        raise FileNotFoundError("未找到 paper_artifacts 下的 RT_A3C_MLN 实验章节 Word 文档。")
    return matches[0]


def set_east_asian_font(run, font="宋体", size=None, bold=None, color=None):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="B7C7D6", size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_section_portrait(section):
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)


def set_section_landscape(section):
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21)
    section.top_margin = Cm(1.55)
    section.bottom_margin = Cm(1.45)
    section.left_margin = Cm(1.35)
    section.right_margin = Cm(1.35)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_east_asian_font(r, font="微软雅黑", size=15 if level == 1 else 12.5, bold=True, color=ACCENT)
    return p


def add_body(doc, text, first_line=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(5)
    if first_line:
        p.paragraph_format.first_line_indent = Pt(21)
    r = p.add_run(text)
    set_east_asian_font(r, font="宋体", size=10.5)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(text)
    set_east_asian_font(r, font="宋体", size=9.2, bold=True, color="333333")
    return p


def add_note_box(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_borders(table, color="BFD3E3", size="6")
    cell = table.cell(0, 0)
    shade_cell(cell, NOTE_FILL)
    set_cell_margins(cell, top=150, bottom=150, start=180, end=180)
    p = cell.paragraphs[0]
    r = p.add_run(title)
    set_east_asian_font(r, font="微软雅黑", size=10.5, bold=True, color=ACCENT)
    p2 = cell.add_paragraph()
    p2.paragraph_format.line_spacing = 1.18
    r2 = p2.add_run(body)
    set_east_asian_font(r2, font="宋体", size=10)
    doc.add_paragraph()


def format_word_table(table, font_size=7.1, header_size=6.9):
    table.style = "Table Grid"
    table.autofit = True
    set_table_borders(table)
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=70, bottom=70, start=55, end=55)
            if i == 0:
                shade_cell(cell, HEADER_FILL)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.line_spacing = 1.03
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    set_east_asian_font(
                        r,
                        font="宋体",
                        size=header_size if i == 0 else font_size,
                        bold=True if i == 0 else None,
                    )
    set_repeat_table_header(table.rows[0])


def add_dataframe_table(doc, df, caption):
    add_caption(doc, caption)
    table = doc.add_table(rows=1, cols=len(df.columns))
    for j, col in enumerate(df.columns):
        table.cell(0, j).text = str(col)
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for j, col in enumerate(df.columns):
            cells[j].text = str(row[col])
    format_word_table(table)
    doc.add_paragraph()


def prepare_table3():
    df = pd.read_csv(TABLES / "Table3_Strong_Rule_Drift.csv")
    rename = {
        "Method": "方法",
        "Adaptive structure": "结构自适应",
        "Post-drift F1": "漂移后F1",
        "Post-drift AUC": "漂移后AUC",
        "Post-drift Log-loss": "漂移后Log-loss",
        "Recovery lag chunks": "恢复延迟(chunk)",
        "Blocking cycle ms": "阻塞周期(ms)",
        "10ms success %": "10ms成功率(%)",
        "Struct update %": "结构更新(%)",
        "Rule +/- per chunk": "规则增/删",
        "Active emerging rules": "激活新规则",
    }
    df = df.rename(columns=rename)
    df["结构自适应"] = df["结构自适应"].replace({"Yes": "是", "No": "否"})
    df["方法"] = df["方法"].replace(
        {
            "OnlineWeight fixed": "OnlineWeight fixed",
            "Static fixed": "Static fixed",
            "Rolling Boosted": "Rolling Boosted",
            "Rolling MaxEnt": "Rolling MaxEnt",
        }
    )
    return df


def add_picture(doc, image, caption, width_inches):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image), width=Inches(width_inches))
    add_caption(doc, caption)


def append_section(doc):
    table3 = prepare_table3()

    doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_portrait(doc.sections[-1])
    add_heading(doc, "4.9 强规则漂移压力测试", 1)
    add_body(
        doc,
        "前述常规数据流实验主要模拟窗口级别的类别先验变化。在这类温和漂移下，OnlineWeight MLN 和 Static MaxEnt reference 由于初始规则结构仍能覆盖后续样本，可能在 F1 或响应时间上表现较好。为了进一步检验固定结构方法的局限性，本节构造强规则漂移压力测试，使后半段数据流中出现初始训练阶段几乎不可见的新判别规则，从而验证规则结构自适应是否必要。",
    )
    add_body(
        doc,
        "该压力测试包含 10 个数据流 chunk，每个 chunk 含 1000 个样本，第 5 个 chunk 为规则漂移点。漂移前标签主要由 old rules 决定；漂移后 emerging rules 的支持度显著升高，并成为主要判别规则。OnlineWeight fixed 和 Static fixed 只允许使用初始阶段选出的固定规则结构，因此即使 OnlineWeight 可以更新已有规则权重，也无法添加后半段才出现的 emerging rules。",
    )

    doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_landscape(doc.sections[-1])
    add_dataframe_table(doc, table3, "表4-3 强规则漂移压力测试结果")
    add_body(
        doc,
        "表4-3 显示，在强规则漂移后，RT-A3C-MLN 的 F1 为 0.704，阻塞周期为 7.871 ms，10 ms 成功率为 100%，并且结构更新比例达到 60.0%，每个 chunk 平均发生 1.20 次规则新增和 1.20 次规则删除。相比之下，OnlineWeight fixed 和 Static fixed 的漂移后 F1 分别下降到 0.344 和 0.326，且均未恢复到设定阈值。这说明固定结构方法虽然具有较低响应时间，但在新规则出现时缺乏结构扩展能力。",
    )
    add_body(
        doc,
        "Rolling Boosted 的漂移后 F1 为 0.707，略高于 RT-A3C-MLN，但其阻塞周期达到 855.713 ms，10 ms 成功率为 0%，无法支撑实时推理结论。Rolling L1 和 Rolling MaxEnt 虽然能够通过滚动重训适应部分漂移，但阻塞周期分别为 40.979 ms 和 23.702 ms，同样无法满足 10 ms 实时约束。因此，强漂移实验进一步说明：本文方法的优势不是单一最高准确率，而是在规则结构可变的同时保持实时响应。",
    )

    doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_portrait(doc.sections[-1])
    add_picture(
        doc,
        FIGURES / "Fig6_Strong_Rule_Drift_Trajectory.png",
        "图4-6 强规则漂移下的恢复轨迹与实时响应",
        width_inches=6.25,
    )
    add_body(
        doc,
        "图4-6 展示了强规则漂移过程中的动态变化。上方子图以 F1 衡量各方法在每个 chunk 上的分类性能，灰色背景表示 emerging rules 的支持度变化，竖向虚线标出第 5 个 chunk 的规则漂移点。可以看到，OnlineWeight fixed 和 Static fixed 在漂移后性能迅速下降，说明固定结构无法表达新出现的判别关系。RT-A3C-MLN 在漂移后经过少量 chunk 即恢复到较高 F1，表明后台规则结构刷新能够吸收新的局部反馈。",
    )
    add_body(
        doc,
        "中间子图显示 Log-loss 的变化。OnlineWeight fixed 在漂移后 Log-loss 明显升高，说明它不仅分类结果变差，概率输出也严重失准。下方子图显示阻塞周期，RT-A3C-MLN 始终低于 10 ms 实时截止线；而 Boosted 和 BeamSearch 类方法虽然具备更重的重训能力，但阻塞周期达到百毫秒甚至秒级，不能作为实时系统的可行方案。",
    )

    add_picture(
        doc,
        FIGURES / "Fig7_Strong_Drift_Structure_Evidence.png",
        "图4-7 强规则漂移下的结构自适应证据",
        width_inches=6.25,
    )
    add_body(
        doc,
        "图4-7 从三个角度总结压力测试结果。左侧子图比较漂移后的 F1，RT-A3C-MLN 明显优于 OnlineWeight fixed 和 Static fixed；中间子图比较漂移后的 Log-loss，OnlineWeight fixed 的概率损失最高，说明固定结构在线权重更新难以提供可靠概率推理；右侧子图显示 RT-A3C-MLN 在漂移点后激活的 emerging rules 数量明显增加，证明模型确实发生了规则结构层面的调整，而不是仅仅调整已有规则权重。",
    )
    add_note_box(
        doc,
        "强漂移实验结论",
        "该压力测试补充说明了常规实验中 OnlineWeight MLN 和 Static MaxEnt 指标较高的原因：当窗口变化较温和时，固定结构仍可能覆盖后续样本；但当新判别规则在后半段才出现时，固定结构方法无法新增规则，性能和概率质量显著下降。RT-A3C-MLN 通过异步结构刷新激活 emerging rules，并保持 10 ms 内阻塞周期，从而更直接地支撑“MLN 规则可以在不同数据流输入下自适应调整并实时推理”的核心论点。",
    )


def main():
    doc_path = target_docx()
    doc = Document(str(doc_path))
    if any("强规则漂移压力测试" in p.text for p in doc.paragraphs):
        raise RuntimeError("目标文档中已存在“强规则漂移压力测试”小节，为避免重复追加已停止。")

    backup = doc_path.with_name(doc_path.stem + "_backup_before_Table3_Fig6_Fig7.docx")
    if not backup.exists():
        shutil.copy2(doc_path, backup)

    append_section(doc)
    doc.save(str(doc_path))
    print(doc_path)


if __name__ == "__main__":
    main()
