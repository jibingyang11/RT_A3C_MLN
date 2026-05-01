from __future__ import annotations

from pathlib import Path
import shutil

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
ARTIFACTS = ROOT / "paper_artifacts"
LATEST = ARTIFACTS / "latest_foundation_baselines"
TABLES = LATEST / "tables"
FIGURES = LATEST / "figures"
DOCX = ARTIFACTS / "RT_A3C_MLN_实验章节报告.docx"

ACCENT = "1F4E79"
HEADER_FILL = "D9EAF7"
NOTE_FILL = "EAF3F8"


def set_east_asian_font(run, font="宋体", size=None, bold=None, color=None):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=60, start=45, bottom=60, end=45):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="B7C7D6", size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_section_portrait(section):
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)


def set_section_landscape(section):
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21)
    section.top_margin = Cm(1.45)
    section.bottom_margin = Cm(1.35)
    section.left_margin = Cm(1.15)
    section.right_margin = Cm(1.15)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_east_asian_font(r, font="微软雅黑", size=15 if level == 1 else 12.5, bold=True, color=ACCENT)
    return p


def add_body(doc, text, first_line=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(5)
    if first_line:
        p.paragraph_format.first_line_indent = Pt(21)
    r = p.add_run(text)
    set_east_asian_font(r, font="宋体", size=10.5)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(text)
    set_east_asian_font(r, font="宋体", size=9.2, bold=True, color="333333")
    return p


def add_note_box(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_borders(table, color="BFD3E3", size="6")
    cell = table.cell(0, 0)
    shade_cell(cell, NOTE_FILL)
    set_cell_margins(cell, top=150, bottom=150, start=180, end=180)
    p = cell.paragraphs[0]
    r = p.add_run(title)
    set_east_asian_font(r, font="微软雅黑", size=10.5, bold=True, color=ACCENT)
    p2 = cell.add_paragraph()
    p2.paragraph_format.line_spacing = 1.18
    r2 = p2.add_run(body)
    set_east_asian_font(r2, font="宋体", size=10)
    doc.add_paragraph()


def format_word_table(table, font_size=5.9, header_size=5.8):
    table.style = "Table Grid"
    table.autofit = True
    set_table_borders(table)
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if i == 0:
                shade_cell(cell, HEADER_FILL)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    set_east_asian_font(
                        r,
                        font="宋体",
                        size=header_size if i == 0 else font_size,
                        bold=True if i == 0 else None,
                    )
    set_repeat_table_header(table.rows[0])


def prepare_table6():
    df = pd.read_csv(TABLES / "Table6_Latest_Foundation_Baselines.csv")
    rename = {
        "Method": "方法",
        "Year/type": "年份/类型",
        "Stream adaptation": "流式适应方式",
        "Rule-structure adaptation": "规则结构自适应",
        "Core F1": "常规F1",
        "Core Log-loss": "常规Log-loss",
        "Core cycle ms": "常规周期(ms)",
        "Core 10ms %": "10ms成功率(%)",
        "Strong post-F1": "强漂移F1",
        "Strong post Log-loss": "强漂移Log-loss",
        "Recovery lag": "恢复延迟",
        "Struct update %": "结构更新(%)",
        "Active emerging rules": "激活新规则",
    }
    df = df.rename(columns=rename)
    df["规则结构自适应"] = df["规则结构自适应"].replace({"Yes": "是", "No": "否"})
    df["年份/类型"] = df["年份/类型"].replace(
        {
            "2026 / ours": "2026/本文",
            "2025 foundation model": "2025基础模型",
            "2026 checkpoint / 2025 paper": "2026 ckpt/2025",
            "2025 / 2024 preprint": "2025/2024",
            "fixed-rule online reference": "固定规则在线参考",
            "static reference": "静态参考",
        }
    )
    df["流式适应方式"] = df["流式适应方式"].replace(
        {
            "A3C rule/weight update": "A3C规则/权重",
            "rolling context": "滚动上下文",
            "online fine-tuning": "在线微调",
            "fixed reference": "固定参考",
        }
    )
    return df


def add_dataframe_table(doc, df, caption):
    add_caption(doc, caption)
    table = doc.add_table(rows=1, cols=len(df.columns))
    for j, col in enumerate(df.columns):
        table.cell(0, j).text = str(col)
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for j, col in enumerate(df.columns):
            cells[j].text = str(row[col])
    format_word_table(table)
    doc.add_paragraph()


def add_picture(doc, image, caption, width_inches=6.35):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image), width=Inches(width_inches))
    add_caption(doc, caption)


def append_section(doc):
    table6 = prepare_table6()

    doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_portrait(doc.sections[-1])
    add_heading(doc, "4.12 2024-2026 最新表格基础模型补充实验", 1)
    add_body(
        doc,
        "为进一步回应对比基线年份偏早的问题，本文在现代数据流基线之外，继续加入 2024-2026 年间具有代表性的表格基础模型和深度表格模型。新增方法包括 TabPFN v2、TabICL v2、TabDPT 和 TabM。其中 TabPFN v2、TabICL v2 和 TabDPT 属于预训练表格基础模型，实验中采用滚动上下文方式吸收最新 chunk；TabM 属于参数高效深度表格集成模型，实验中采用轻量在线微调方式适应数据流。",
    )
    add_body(
        doc,
        "这些方法的加入可以提高对比基线的新近性，但也需要明确其与本文方法的差异：表格基础模型主要通过上下文推理或神经网络参数更新获得预测能力，并不执行 MLN 规则的新增、删除、结构替换或规则权重解释。因此，该补充实验的目的不是证明本文方法在所有静态预测指标上超过最新表格模型，而是检验在更强的新近基线面前，本文方法是否仍能保持实时响应并提供规则结构自适应证据。",
    )

    doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_landscape(doc.sections[-1])
    add_dataframe_table(doc, table6, "表4-6 2024-2026 最新表格基础模型补充实验结果")
    add_body(
        doc,
        "表4-6 显示，在常规数据流实验中，TabPFN v2 的 F1 为 0.695，略高于 RT-A3C-MLN 的 0.689；TabICL v2 的 F1 为 0.662，也具有一定竞争力。这说明最新表格基础模型确实是强预测基线，应当作为论文中的近年补充对比。然而，从实时性角度看，TabPFN v2、TabICL v2 和 TabDPT 的平均阻塞周期分别为 4281.724 ms、2116.009 ms 和 7771.539 ms，均远高于 10 ms 实时阈值；TabM 的阻塞周期为 13.419 ms，也未满足 10 ms 截止约束。相比之下，RT-A3C-MLN 的阻塞周期为 2.340 ms，10 ms 成功率为 100%。",
    )
    add_body(
        doc,
        "在强规则漂移场景下，TabPFN v2、TabICL v2 和 OnlineWeight MLN 的 post-F1 分别为 0.452、0.488 和 0.497，虽然在局部分类指标上接近本文方法，但它们的结构更新比例和激活新规则数量均为 0，不能说明 MLN 规则结构发生了自适应调整。RT-A3C-MLN 的结构更新比例为 75.0%，并激活 1.25 个 emerging rules，说明模型在数据流后半段确实进行了规则层面的结构更新。该结果强化了本文的核心论点：最新强预测器可以作为准确率参照，但不能替代概率逻辑规则网络的实时结构自适应机制。",
    )

    doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_portrait(doc.sections[-1])
    add_picture(
        doc,
        FIGURES / "Fig10_Latest_Foundation_Baselines.png",
        "图4-10 2024-2026 最新表格基础模型与本文方法的对比",
        width_inches=6.35,
    )
    add_body(
        doc,
        "图4-10 左侧子图展示各方法在常规数据流中的阻塞周期，纵轴采用对数坐标，红色虚线表示 10 ms 实时阈值。可以看到，最新表格基础模型的预测能力较强，但上下文推理或在线微调带来的阻塞成本明显高于实时阈值。中间子图比较常规数据流 F1 和强规则漂移后的 post-F1，说明 TabPFN v2 和 TabICL v2 能提供较强的分类参考，但在强漂移场景中仍未表现出规则结构恢复机制。",
    )
    add_body(
        doc,
        "右侧子图展示 MLN 规则结构证据。只有 RT-A3C-MLN 出现非零的结构更新比例和 active emerging rules，说明本文方法在数据流中不仅更新预测输出，也更新概率逻辑知识网络的规则结构。TabPFN v2、TabICL v2、TabDPT 和 TabM 的对应指标均为 0，表明它们虽然属于更新近的强基线，但并不提供本文所需的规则新增、删除和结构调整证据。",
    )
    add_note_box(
        doc,
        "最新基线补充实验结论",
        "2024-2026 最新表格模型显著增强了论文对比的时效性。实验表明，TabPFN v2、TabICL v2、TabDPT 和 TabM 可作为强预测参照，但在 CPU 实时数据流场景下阻塞周期较高，且不具备 MLN 规则结构自适应能力。RT-A3C-MLN 的主要贡献仍然成立：在保持低阻塞周期的同时，通过异步 Actor-Critic 机制实现规则结构和权重的可解释自适应。",
    )


def main():
    if not DOCX.exists():
        raise FileNotFoundError(DOCX)
    doc = Document(str(DOCX))
    if any("2024-2026 最新表格基础模型补充实验" in paragraph.text for paragraph in doc.paragraphs):
        raise RuntimeError("目标文档中已存在最新表格基础模型补充实验小节，已停止以避免重复追加。")
    backup = DOCX.with_name(DOCX.stem + "_backup_before_Table6_Fig10.docx")
    if not backup.exists():
        shutil.copy2(DOCX, backup)
    append_section(doc)
    doc.save(str(DOCX))
    print(DOCX)


if __name__ == "__main__":
    main()
