"""Build the formal Experiments chapter DOCX from generated CSV tables and figures.

This script is intentionally separated from the experiment runner so the report can
be rebuilt after visual/layout tweaks without rerunning all baselines.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
ARTIFACT_DIR = ROOT / "paper_artifacts"
TABLE_DIR = ARTIFACT_DIR / "tables"
FIGURE_DIR = ARTIFACT_DIR / "figures"
OUT_DOCX = ARTIFACT_DIR / "RT_A3C_MLN_正式实验章节.docx"


ACCENT = "1F5E5B"
LIGHT = "E9F2F0"
HEADER = "DCEBE7"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=55, start=55, bottom=55, end=55) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_font(paragraph, size: float = 10.5, bold: bool = False) -> None:
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(level=level)
    run = p.add_run(text)
    set_run_font(run, size=16 if level == 1 else 13, bold=True, color=ACCENT)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    set_paragraph_font(p, 10.5)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=9, bold=True, color=ACCENT)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)


def add_figure(doc: Document, filename: str, caption: str, width_cm: float = 25.2) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(FIGURE_DIR / filename), width=Cm(width_cm))
    add_caption(doc, caption)


def make_table(doc: Document, df: pd.DataFrame, caption: str, font_size: float = 5.4) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(caption)
    set_run_font(run, size=9, bold=True, color=ACCENT)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)

    table = doc.add_table(rows=1, cols=len(df.columns))
    table.style = "Table Grid"
    table.autofit = True
    for i, col in enumerate(df.columns):
        cell = table.rows[0].cells[i]
        cell.text = str(col)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, HEADER)
        set_cell_margins(cell, 45, 45, 45, 45)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_font(paragraph, font_size, True)

    for _, row in df.iterrows():
        cells = table.add_row().cells
        for i, col in enumerate(df.columns):
            value = row[col]
            text = "" if pd.isna(value) else str(value)
            cells[i].text = text
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cells[i], 35, 40, 35, 40)
            if row.name % 2 == 1:
                set_cell_shading(cells[i], "F7FAFA")
            if str(row.get("Method", "")) == "RT-A3C-MLN" or str(row.get("Variant", "")) == "Full RT-A3C":
                set_cell_shading(cells[i], LIGHT)
            for paragraph in cells[i].paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if i in (0, 1, 2) else WD_ALIGN_PARAGRAPH.CENTER
                set_paragraph_font(paragraph, font_size, i == 0)

    doc.add_paragraph()


def compact_main_table(df: pd.DataFrame) -> pd.DataFrame:
    cols = [
        "Method",
        "Group",
        "Year/type",
        "Rule adapt",
        "Core F1",
        "Core cycle ms",
        "10ms %",
        "Strong F1",
        "Strong Log-loss",
        "Recovery",
        "Struct %",
        "Rule +/-",
        "Emerging rules",
    ]
    return df[cols].copy()


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.1)
    section.right_margin = Cm(1.1)

    styles = doc.styles
    styles["Normal"].font.name = "宋体"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    styles["Normal"].font.size = Pt(10.5)


def main() -> None:
    main_table = pd.read_csv(TABLE_DIR / "Table1_Formal_All_Methods.csv")
    ablation_table = pd.read_csv(TABLE_DIR / "Table2_Formal_Ablation.csv")

    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("实验章节：基于异步优势 Actor-Critic 的 MLN 实时规则学习")
    set_run_font(run, size=18, bold=True, color=ACCENT)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("正式实验结果、强规则漂移压力测试与消融分析")
    set_run_font(run, size=10.5, color="4A4A4A")

    add_heading(doc, "实验目标", 1)
    add_body(
        doc,
        "本章围绕“MLN 规则可以在不同数据流输入下自适应调整，并形成能够实时响应输入数据、提供推理结果的概率知识网络”这一核心论点展开。实验不把准确率作为唯一目标，而同时考察预测质量、推理/更新周期、10ms 实时达标率、强规则漂移下的恢复速度，以及规则结构是否发生添加、删除和权重更新。"
    )
    add_body(
        doc,
        "对比方法覆盖传统 MLN 结构学习、固定规则在线权重更新、数据流学习器、近年表格基础模型，以及近年规则自适应方法的 style/proxy 复现。style/proxy 表示在统一的规则特征和数据流协议下复现其核心机制，用于比较实时性和规则自适应行为，并不声称调用原作者官方实现。"
    )

    add_figure(doc, "Fig1_Method_Framework.png", "图1  RT-A3C-MLN 方法框架。", 23.5)

    add_heading(doc, "主结果", 1)
    add_body(
        doc,
        "表1汇总所有对比方法。Core F1 表示常规真实数据流上的平均分类质量，Core cycle ms 表示每个数据块完成推理与在线更新的平均周期时间，10ms % 表示周期时间低于 10ms 的数据块比例。Strong F1 和 Strong Log-loss 来自强规则漂移压力测试，用于观察规则发生剧烈变化时模型是否仍能保持有效概率推理。"
    )
    make_table(doc, compact_main_table(main_table), "表1  所有对比方法的正式主结果。", 5.2)
    add_figure(doc, "Fig2_All_Method_Dashboard.png", "图2  所有方法的实时性、预测质量、漂移质量与规则结构证据总览。", 25.0)

    add_heading(doc, "强规则漂移压力测试", 1)
    add_body(
        doc,
        "图3展示强规则漂移过程中各方法在不同 chunk 上的表现。该实验人为设置规则主导关系随时间显著变化，用于检验固定规则 MLN 和只更新权重的方法是否会出现更新滞后。RT-A3C-MLN 的关键优势不在于始终取得最高静态 F1，而在于可以在低周期时间内持续刷新规则结构，并在漂移后较快恢复。"
    )
    add_figure(doc, "Fig3_Strong_Drift_Heatmap.png", "图3  强规则漂移下的逐 chunk 热力图。", 25.0)

    add_heading(doc, "消融实验", 1)
    add_body(
        doc,
        "消融实验考察并行 Actor、先验漂移修正、规则结构刷新和仅在线权重更新对方法的贡献。其中 No structure refresh 与 Online only 不允许添加、删除或替换规则，因此它们在强规则漂移下缺少结构层面的自适应能力。"
    )
    make_table(doc, ablation_table, "表2  RT-A3C-MLN 关键组件消融结果。", 7.2)
    add_figure(doc, "Fig4_Ablation_Rule_Evidence.png", "图4  消融实验中的规则结构证据。", 21.2)

    add_heading(doc, "结果结论", 1)
    add_body(
        doc,
        "综合表1和图2可见，固定规则方法虽然单次更新很快，但在强漂移数据流中无法产生规则结构变化，概率质量和恢复能力明显不足；表格基础模型通常具有较强预测能力，但推理周期远高于实时响应要求；传统 MLN 搜索类方法能表达规则结构，却难以维持低延迟。"
    )
    add_body(
        doc,
        "因此，RT-A3C-MLN 的证据链不是单一 F1，而是低周期时间、10ms 达标、漂移恢复和规则结构更新共同支持实时规则自适应。"
    )

    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    main()
